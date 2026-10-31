#!/usr/bin/env python3
"""
Turns a content-plan workbook (Stacko Content Studio .xlsx) into one SQL file
that creates a project per planned post (sheet "Content") and per backlog idea
(sheet "Backlog") through public.create_project() — the same RPC every intake
channel uses (P4-T10).

- Idempotent: each row has a source key (content-plan:<workbook>:<ID>); running
  the SQL again skips rows that already exist.
- Rows land as status `new` (auto-analysis off), tagged, with the full plan row
  in source_metadata.content_plan so later analysis can compare new ideas with
  existing content and avoid duplicates.
- Content already posted is closed later from the CRM (Idea tab → "Mark as
  published" with the post's link).

Usage (needs openpyxl):
  python3 -I scripts/import/content-plan-to-sql.py <workbook.xlsx> <out.sql> [--email you@example.com] [--brand Stacko]
"""
import argparse
import datetime as dt
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

WORKBOOK_KEY = "stacko-content-studio"
MAX_IDEA = 19_000


def normalize_idea(text: str) -> str:
    """Same as normalizeIdeaText() in packages/shared/src/hash.ts."""
    t = unicodedata.normalize("NFKC", text).lower()
    t = re.sub("[ً-ٰٟـ]", "", t)
    return re.sub(r"\s+", " ", t).strip()


def dedupe_hash(idea: str) -> str:
    """Same as dedupeHash(idea, null): sha256 of "<normalized idea>\\n"."""
    return hashlib.sha256((normalize_idea(idea) + "\n").encode("utf-8")).hexdigest()


def clean(v):
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def slug(s: str, limit: int = 50) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
    return s[:limit].rstrip("-")


def content_type(fmt: str | None) -> str | None:
    f = (fmt or "").lower()
    if "carousel" in f or "stor" in f or "poll" in f or "quote card" in f:
        return "carousel"
    if "reel" in f or "video" in f or "live" in f:
        return "video"
    if "static" in f or "post" in f or "meme" in f or "one-pager" in f:
        return "static"
    return None


def format_tag(fmt: str | None) -> str | None:
    f = (fmt or "").lower()
    for word in ("reel", "carousel", "stories", "story", "static"):
        if word in f:
            return "stories" if word == "story" else word
    return None


def sheet_rows(ws, header_row: int):
    headers = [clean(c) for c in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or not clean(row[0]):
            continue
        yield {h: clean(v) for h, v in zip(headers, row) if h and clean(v) is not None}


def compose(parts: list[tuple[str, str | None]]) -> str:
    text = "\n\n".join(f"{label}:\n{value}" if label else value for label, value in parts if value)
    return text[:MAX_IDEA]


def content_record(r: dict, source_file: str) -> dict:
    pid = r["ID"]
    fmt = r.get("Format")
    head = f"[{pid}] {r.get('Title', '')}\n" + " · ".join(
        str(x) for x in (fmt, r.get("Duration / slides"), r.get("Pillar"), r.get("Funnel"), r.get("Publish date")) if x
    )
    idea = compose(
        [
            ("", head),
            ("Objective", r.get("Objective")),
            ("Hook (EN)", r.get("Hook EN (on screen, first 1.5 s)")),
            ("Hook (AR)", r.get("Hook AR")),
            ("Scene plan", r.get("Scene plan (beat by beat)")),
            ("Script", r.get("Original script (from the plan)")),
            ("Caption", r.get("Caption")),
            ("Hashtags", r.get("Hashtags")),
            ("CTA", r.get("CTA")),
        ]
    )
    month = (r.get("Publish date") or "")[:7]
    tags = ["content-plan", f"plan-{month}" if month else None, format_tag(fmt), slug(r.get("Pillar") or "") or None]
    return {
        "source_key": f"content-plan:{WORKBOOK_KEY}:{pid}",
        "dedupe_hash": dedupe_hash(idea),
        "tags": [t for t in dict.fromkeys(tags) if t],
        "record": {
            "title": f"{pid} · {r.get('Title', pid)}"[:200],
            "original_idea": idea,
            "input_source": "csv",
            "source_reference": f"{source_file} › Content › {pid}",
            "content_type": content_type(fmt),
            "language": "bilingual",
            "source_metadata": {
                "origin": "content_plan_import",
                "source_file": source_file,
                "sheet": "Content",
                "plan_id": pid,
                "plan_status": r.get("Status"),
                "publish_date": r.get("Publish date"),
                "format": fmt,
                "content_folder": r.get("Content folder"),
                "content_plan": r,
            },
        },
    }


def backlog_record(r: dict, source_file: str) -> dict:
    pid = r["ID"]
    fmt = r.get("Content type")
    head = f"[{pid}] {r.get('Title', '')}\n" + " · ".join(
        str(x) for x in (fmt, r.get("Pillar"), r.get("Funnel"), r.get("Batch")) if x
    )
    idea = compose(
        [
            ("", head),
            ("Objective", r.get("Objective")),
            ("Hook", r.get("Hook")),
            ("Script", r.get("Full script")),
            ("Caption", r.get("Caption")),
            ("CTA", r.get("CTA")),
            ("Overlaps with", r.get("Overlaps with")),
        ]
    )
    batch = re.match(r"Batch (\d+)", r.get("Batch") or "")
    tags = ["content-plan", "backlog", f"backlog-batch-{batch.group(1)}" if batch else None, format_tag(fmt), slug(r.get("Pillar") or "") or None]
    return {
        "source_key": f"content-plan:{WORKBOOK_KEY}:{pid}",
        "dedupe_hash": dedupe_hash(idea),
        "tags": [t for t in dict.fromkeys(tags) if t],
        "record": {
            "title": f"{pid} · {r.get('Title', pid)}"[:200],
            "original_idea": idea,
            "input_source": "csv",
            "source_reference": f"{source_file} › Backlog › {pid}",
            "content_type": content_type(fmt),
            "language": "bilingual",
            "source_metadata": {
                "origin": "content_plan_import",
                "source_file": source_file,
                "sheet": "Backlog",
                "plan_id": pid,
                "plan_status": r.get("Status"),
                "publish_date": None,
                "format": fmt,
                "overlaps_with": r.get("Overlaps with"),
                "content_plan": r,
            },
        },
    }


SQL = """-- Import of {source_file} into the content pipeline (generated by scripts/import/content-plan-to-sql.py).
-- {n_content} planned posts (sheet Content) + {n_backlog} backlog ideas (sheet Backlog) → one project each.
--
-- Run it in the Supabase SQL editor (or psql) as the postgres role.
--   • Edit v_email below if you sign in to the CRM with another address.
--   • Safe to run again: rows that already exist are skipped (source key per row ID).
--   • Every row arrives as a new idea (no AI analysis started), tagged "content-plan",
--     with the full plan row kept in source_metadata.content_plan.
--   • When a post is live: open it in the CRM → Idea tab → "Mark as published" + link.
--   • Backlog ideas are tagged "backlog"; to skip them, delete the objects whose
--     "sheet" is "Backlog" from the list below before running.

do $import$
declare
  v_email text := {email};          -- the CRM login that owns the imported rows
  v_brand_name text := {brand};     -- brand to attach, if it exists in the workspace
  v_actor uuid;
  v_ws uuid;
  v_brand uuid;
  v_batch uuid;
  r jsonb;
  res jsonb;
  v_created int := 0;
  v_existing int := 0;
  v_duplicates int := 0;
  v_report jsonb := '[]'::jsonb;
begin
  select id into v_actor from auth.users where lower(email) = lower(v_email);
  if v_actor is null then
    raise exception 'No CRM user with email % — edit v_email at the top of this script', v_email;
  end if;
  select workspace_id into v_ws from public.workspace_members
  where user_id = v_actor and role in ('admin', 'editor') order by created_at limit 1;
  if v_ws is null then
    raise exception 'User % is not an admin or editor of any workspace', v_email;
  end if;
  select id into v_brand from public.brands
  where workspace_id = v_ws and archived_at is null and lower(name) = lower(v_brand_name) limit 1;
  if v_brand is null then
    raise notice 'Brand "%" not found in the workspace — rows are imported without a brand', v_brand_name;
  end if;

  insert into public.import_batches (workspace_id, brand_id, file_name, field_mapping, status, total_rows, created_by)
  values (v_ws, v_brand, {source_file_sql}, '{{"origin": "content_plan_import", "sheets": ["Content", "Backlog"]}}'::jsonb,
          'processing', {n_total}, v_actor)
  returning id into v_batch;

  for r in select value from jsonb_array_elements($plan${rows}$plan$::jsonb) loop
    res := public.create_project(
      v_ws, v_actor,
      (r -> 'record') || jsonb_build_object('brand_id', v_brand),
      r ->> 'dedupe_hash', false, r ->> 'source_key', false, v_batch);
    if (res ->> 'created')::boolean then
      v_created := v_created + 1;
      insert into public.tags (workspace_id, name)
      select v_ws, t from jsonb_array_elements_text(r -> 'tags') t
      on conflict (workspace_id, name) do nothing;
      insert into public.project_tags (project_id, tag_id)
      select (res ->> 'project_id')::uuid, tg.id from public.tags tg
      where tg.workspace_id = v_ws and tg.name in (select jsonb_array_elements_text(r -> 'tags'))
      on conflict do nothing;
    elsif (res ->> 'replayed')::boolean then
      v_existing := v_existing + 1;
    else
      v_duplicates := v_duplicates + 1;
    end if;
    v_report := v_report || jsonb_build_array(jsonb_build_object(
      'plan_id', r -> 'record' -> 'source_metadata' ->> 'plan_id',
      'project_id', res ->> 'project_id',
      'result', case when (res ->> 'created')::boolean then 'created'
                     when (res ->> 'replayed')::boolean then 'already_imported' else 'duplicate' end));
  end loop;

  update public.import_batches
  set status = 'completed', created_count = v_created, duplicate_count = v_duplicates, report = v_report
  where id = v_batch;
  raise notice 'Content plan import: % created, % already imported, % duplicates of existing ideas (batch %)',
    v_created, v_existing, v_duplicates, v_batch;
end
$import$;
"""


def sql_literal(s: str) -> str:
    return "'" + s.replace("'", "''") + "'"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("out")
    ap.add_argument("--email", default="you@example.com")
    ap.add_argument("--brand", default="Stacko")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=True, read_only=True)
    source_file = Path(args.workbook).name
    source_file = re.sub(r"^[0-9a-f]{8}-", "", source_file)  # drop upload prefixes
    content = [content_record(r, source_file) for r in sheet_rows(wb["Content"], 2)]
    backlog = [backlog_record(r, source_file) for r in sheet_rows(wb["Backlog"], 1)] if "Backlog" in wb.sheetnames else []
    rows = content + backlog
    ids = [r["record"]["source_metadata"]["plan_id"] for r in rows]
    if len(ids) != len(set(ids)):
        print("duplicate IDs in the workbook", file=sys.stderr)
        return 1
    payload = json.dumps(rows, ensure_ascii=False, indent=1)
    if "$plan$" in payload or "$import$" in payload:
        print("content contains a dollar-quote tag", file=sys.stderr)
        return 1
    Path(args.out).write_text(
        SQL.format(
            source_file=source_file,
            source_file_sql=sql_literal(source_file),
            email=sql_literal(args.email),
            brand=sql_literal(args.brand),
            n_content=len(content),
            n_backlog=len(backlog),
            n_total=len(rows),
            rows=payload,
        ),
        encoding="utf-8",
    )
    print(f"wrote {args.out}: {len(content)} planned posts + {len(backlog)} backlog ideas")
    return 0


if __name__ == "__main__":
    sys.exit(main())
